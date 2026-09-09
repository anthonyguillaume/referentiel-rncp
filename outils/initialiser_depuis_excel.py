#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Réinitialise state/snapshot.json à partir d'un fichier Excel de liens
formacode → RNCP (colonne « Nomenclature RNCP » ou « Code RNCP »).

Les codes du fichier deviennent la base de référence, considérés comme
ACTIFS et SANS SUCCESSEUR au moment où le fichier a été établi ; les autres
champs (intitulé, type, niveau, échéance, formacodes) sont repris de l'export
courant. Les générations suivantes signalent en orange les fiches devenues
inactives ou remplacées depuis, en vert les diplômes du périmètre absents du
fichier, en rouge les fiches disparues — et ces couleurs persistent jusqu'à
la prochaine réinitialisation. Le journal des générations est conservé.

Usage :
  python outils/initialiser_depuis_excel.py "Liens CPF.xlsx" [--xml export.xml]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import tempfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import requests  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from genere_rncp import canon, log, parser_export, telecharger_et_extraire  # noqa: E402


def codes_du_fichier(chemin: Path) -> set[str]:
    codes: set[str] = set()
    wb = load_workbook(chemin, read_only=True, data_only=True)
    for ws in wb.worksheets:
        for ligne in ws.iter_rows(values_only=True):
            for v in (ligne or [])[:6]:
                if isinstance(v, str) and re.search(r"RNCP\s*\d+", v, re.I):
                    c = canon(v)
                    if c:
                        codes.add(c)
    return codes


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    p.add_argument("excel")
    p.add_argument("--xml", default="")
    p.add_argument("--state", default="state/snapshot.json")
    args = p.parse_args()

    codes = codes_du_fichier(Path(args.excel))
    log(f"{len(codes)} codes RNCP distincts dans {Path(args.excel).name}")

    with tempfile.TemporaryDirectory() as tmp:
        if args.xml:
            xml_path, source = Path(args.xml), Path(args.xml).name
        else:
            with requests.Session() as session:
                xml_path, source = telecharger_et_extraire(session, Path(tmp))
        repertoire = parser_export(xml_path)

    jour = datetime.now().date().isoformat()
    etat, absents = {}, []
    for c in sorted(codes, key=lambda x: int(x[4:])):
        f = repertoire.get(c)
        if not f:
            absents.append(c)
            continue
        etat[c] = {"intitule": f["intitule"], "type": f["type"], "niveau": f["niveau"],
                   "actif": True, "date_fin": f["date_fin"], "successeurs": [],
                   "formacodes": f["formacodes"], "premiere_vue": jour,
                   "dernier_changement": "", "evolution": "", "detail": ""}
    if absents:
        log(f"  {len(absents)} code(s) du fichier absent(s) de l'export, ignoré(s) : "
            + ", ".join(absents[:20]))

    chemin = Path(args.state)
    journal: list = []
    if chemin.exists():
        try:
            precedent = json.loads(chemin.read_text(encoding="utf-8"))
            if precedent.get("version") == 2:
                journal = precedent.get("journal", [])  # l'historique des runs est conservé
        except ValueError:
            pass
    chemin.parent.mkdir(parents=True, exist_ok=True)
    chemin.write_text(json.dumps({
        "version": 2,
        "derniere_execution": datetime.now().isoformat(),
        "source_export": f"{Path(args.excel).name} (référence initiale, export {source})",
        "reference": {"nom": Path(args.excel).name, "date": jour},
        "codes": etat,
        "historique": [{"date": jour, "code": "—", "type": "Initialisation du suivi",
                        "avant": "", "apres": f"{len(etat)} diplômes du fichier {Path(args.excel).name}"}],
        "journal": journal,
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    log(f"Base de référence écrite : {chemin} ({len(etat)} codes)")


if __name__ == "__main__":
    main()

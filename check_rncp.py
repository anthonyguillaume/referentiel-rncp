#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Contrôle des fiches RNCP du fichier "Liens CPF par formacode et diplômes"
contre l'export open data officiel de France compétences (data.gouv.fr).

Produit :
  - une copie annotée du classeur (colonnes Statut / Contrôle FC / codes et
    intitulés à jour, lignes colorées, onglets Synthèse + Changements) ;
  - un résumé texte (summary.md) et HTML (summary.html) pour l'e-mail ;
  - un instantané (state/snapshot.json) pour ne signaler que les nouveautés
    d'une exécution à l'autre.

Statuts :
  VERT   : fiche active, intitulé conforme, pas de successeur publié.
  ORANGE : vigilance — intitulé officiel différent, successeur déjà publié
           alors que la fiche est encore active, ou échéance proche.
  ROUGE  : fiche inactive (remplacée ou expirée sans successeur).
  GRIS   : code introuvable dans le répertoire (saisie à vérifier).

Usage local :
  python check_rncp.py                       # télécharge le dernier export
  python check_rncp.py --xml export.xml      # utilise un export déjà téléchargé
  python check_rncp.py --debug-fiche RNCP35803   # affiche le XML brut d'une fiche
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import tempfile
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import TableColumn

# --------------------------------------------------------------------------- #
# Constantes
# --------------------------------------------------------------------------- #

DATASET_SLUG = (
    "repertoire-national-des-certifications-professionnelles-et-repertoire-specifique"
)
API_V2_RESSOURCES = f"https://www.data.gouv.fr/api/2/datasets/{DATASET_SLUG}/resources/"
API_V1_DATASET = f"https://www.data.gouv.fr/api/1/datasets/{DATASET_SLUG}/"
HTTP_HEADERS = {"User-Agent": "controle-rncp/1.0 (suivi interne des offres CPF)"}

RE_RNCP = re.compile(r"RNCP\s*0*(\d+)", re.IGNORECASE)
RE_DATE_NOM_FICHIER = re.compile(r"(20\d{2})-(\d{2})-(\d{2})")
RE_VERSION = re.compile(r"v(\d+)[-_.](\d+)", re.IGNORECASE)

STATUTS = ("VERT", "ORANGE", "ROUGE", "GRIS")
LIBELLES = {
    "VERT": "RAS",
    "ORANGE": "À vérifier",
    "ROUGE": "Expiré / remplacé",
    "GRIS": "Introuvable",
}
COULEURS = {  # remplissages clairs + couleur de police assortie
    "VERT": ("C6EFCE", "006100"),
    "ORANGE": ("FFEB9C", "9C6500"),
    "ROUGE": ("FFC7CE", "9C0006"),
    "GRIS": ("D9D9D9", "3F3F3F"),
}
FILLS = {
    s: PatternFill(start_color=c[0], end_color=c[0], fill_type="solid")
    for s, c in COULEURS.items()
}
FONTS_STATUT = {s: Font(color=c[1], bold=True) for s, c in COULEURS.items()}

NOUVELLES_COLONNES = [
    "Statut",
    "Contrôle France compétences",
    "Code(s) à jour",
    "Intitulé à jour",
    "Échéance fiche",
    "Formacodes de la fiche à jour",
    "Changé ?",
    "Changement depuis le dernier envoi",
]
FILL_NOUVEAU = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
FONT_NOUVEAU = Font(bold=True, color="1F4E79")


# --------------------------------------------------------------------------- #
# Utilitaires
# --------------------------------------------------------------------------- #

def canon(valeur) -> str | None:
    """Normalise n'importe quelle écriture d'un code RNCP -> 'RNCP35803'."""
    if valeur is None:
        return None
    if isinstance(valeur, (int, float)) and not isinstance(valeur, bool):
        n = int(valeur)
        return f"RNCP{n}" if n > 0 else None
    m = RE_RNCP.search(str(valeur))
    return f"RNCP{int(m.group(1))}" if m else None


def norm_titre(s: str | None) -> str:
    """Normalisation d'intitulé pour comparaison (accents, casse, tirets, espaces)."""
    s = (s or "").replace("œ", "oe").replace("Œ", "OE").replace("æ", "ae").replace("Æ", "AE")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.casefold()
    s = re.sub(r"[’´`]", "'", s)
    s = re.sub(r"[‐‑–—−]", "-", s)
    s = re.sub(r"[\s\u00a0]+", " ", s)
    s = re.sub(r"\s*-\s*", "-", s)
    s = re.sub(r"\s*:\s*", ":", s)
    return s.strip()


def parse_date_souple(s: str | None) -> date | None:
    if not s:
        return None
    s = s.strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def fr_date(d: date | None) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def strip_ns(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def log(msg: str) -> None:
    print(msg, flush=True)


# --------------------------------------------------------------------------- #
# 1. Récupération du dernier export RNCP sur data.gouv.fr
# --------------------------------------------------------------------------- #

def lister_ressources(session: requests.Session) -> list[dict]:
    """Liste les ressources du jeu de données (API v2 paginée, secours v1)."""
    ressources: list[dict] = []
    try:
        for page in (1, 2):
            r = session.get(
                API_V2_RESSOURCES,
                params={"page": page, "page_size": 50},
                headers=HTTP_HEADERS,
                timeout=30,
            )
            r.raise_for_status()
            data = r.json().get("data", [])
            ressources.extend(data)
            if len(data) < 50:
                break
    except Exception as exc:  # noqa: BLE001
        log(f"  API v2 indisponible ({exc}), tentative via l'API v1…")
    if not ressources:
        r = session.get(API_V1_DATASET, headers=HTTP_HEADERS, timeout=60)
        r.raise_for_status()
        ressources = r.json().get("resources", [])
    if not ressources:
        raise SystemExit("Impossible de lister les ressources du jeu de données data.gouv.fr")
    return ressources


def candidats_export_rncp(ressources: list[dict]) -> list[dict]:
    """Filtre et ordonne les zips d'export RNCP (XML), du plus récent au plus ancien."""
    candidats = []
    for res in ressources:
        titre = f"{res.get('title') or ''} {res.get('url') or ''}".lower()
        est_zip = (res.get("format") or "").lower() == "zip" or titre.rstrip().endswith(".zip")
        if not est_zip or "rncp" not in titre or "csv" in titre:
            continue
        m = RE_DATE_NOM_FICHIER.search(titre)
        d = date(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else date.min
        v = RE_VERSION.search(titre)
        version = (int(v.group(1)), int(v.group(2))) if v else (0, 0)
        candidats.append({"res": res, "date": d, "version": version})
    candidats.sort(key=lambda c: (c["date"], c["version"]), reverse=True)
    return [c["res"] for c in candidats]


def telecharger_et_extraire(session: requests.Session, tmpdir: Path) -> tuple[Path, str]:
    """Télécharge le dernier export RNCP et en extrait le XML. -> (chemin_xml, nom_source)."""
    ressources = lister_ressources(session)
    candidats = candidats_export_rncp(ressources)
    if not candidats:
        raise SystemExit(
            "Aucun export RNCP (.zip) trouvé dans le jeu de données — "
            "le nommage des ressources a peut-être changé."
        )
    derniere_erreur = None
    for res in candidats[:3]:
        nom = res.get("title") or res.get("url", "").rsplit("/", 1)[-1]
        url = res.get("url")
        log(f"  Téléchargement : {nom}")
        try:
            zip_path = tmpdir / "export_rncp.zip"
            with session.get(url, headers=HTTP_HEADERS, timeout=(15, 600), stream=True) as r:
                r.raise_for_status()
                with open(zip_path, "wb") as f:
                    shutil.copyfileobj(r.raw, f, length=1024 * 1024)
            taille = zip_path.stat().st_size / 1e6
            log(f"  Reçu : {taille:.1f} Mo")
            with zipfile.ZipFile(zip_path) as z:
                membres = [i for i in z.infolist() if i.filename.lower().endswith(".xml")]
                if not membres:
                    raise ValueError("aucun fichier .xml dans l'archive")
                membre = max(membres, key=lambda i: i.file_size)
                xml_path = tmpdir / "export_rncp.xml"
                with z.open(membre) as src, open(xml_path, "wb") as dst:
                    shutil.copyfileobj(src, dst, length=1024 * 1024)
            log(f"  XML extrait : {membre.filename} ({membre.file_size / 1e6:.0f} Mo)")
            return xml_path, nom
        except Exception as exc:  # noqa: BLE001
            derniere_erreur = exc
            log(f"  Échec sur cette ressource ({exc}), essai suivant…")
    raise SystemExit(f"Téléchargement de l'export RNCP impossible : {derniere_erreur}")


# --------------------------------------------------------------------------- #
# 2. Analyse du XML (streaming, tolérant aux variations de schéma)
# --------------------------------------------------------------------------- #

def _texte(elem: ET.Element, tag: str) -> str | None:
    for enfant in elem:
        if strip_ns(enfant.tag) == tag:
            t = "".join(enfant.itertext()).strip()
            return t or None
    return None


def _codes_lies(elem: ET.Element, motif_tag: str) -> list[str]:
    """Tous les codes RNCP trouvés sous les éléments dont le tag contient motif_tag."""
    codes: set[str] = set()
    for node in elem.iter():
        if motif_tag in strip_ns(node.tag).upper():
            for m in RE_RNCP.finditer("".join(node.itertext())):
                codes.add(f"RNCP{int(m.group(1))}")
    return sorted(codes, key=lambda c: int(c[4:]))


def _formacodes(elem: ET.Element) -> list[tuple[str, str]]:
    resultat, vus = [], set()
    for node in elem.iter():
        if strip_ns(node.tag).upper() == "FORMACODE":
            code = (node.findtext("CODE") or "").strip()
            lib = (node.findtext("LIBELLE") or "").strip()
            if not code:
                m = re.search(r"\b(\d{5})\b", "".join(node.itertext()))
                code = m.group(1) if m else ""
            if code and code not in vus:
                vus.add(code)
                resultat.append((code, lib))
    return resultat


def parser_export(xml_path: Path, debug_fiche: str | None = None) -> dict[str, dict]:
    """Parcourt l'export en streaming et renvoie {code RNCP: infos fiche}."""
    repertoire: dict[str, dict] = {}
    it = ET.iterparse(str(xml_path), events=("start", "end"))
    _, racine = next(it)
    n = 0
    for event, elem in it:
        if event != "end" or strip_ns(elem.tag) != "FICHE":
            continue
        numero = canon(_texte(elem, "NUMERO_FICHE"))
        if numero:
            if debug_fiche and numero == canon(debug_fiche):
                print(ET.tostring(elem, encoding="unicode")[:30000])
                raise SystemExit(0)
            actif_txt = (_texte(elem, "ACTIF") or "").strip().lower()
            abrege_code = abrege_lib = None
            for enfant in elem:
                if strip_ns(enfant.tag) == "ABREGE":
                    abrege_code = (enfant.findtext("CODE") or "").strip() or None
                    abrege_lib = (enfant.findtext("LIBELLE") or "").strip() or None
            repertoire[numero] = {
                "intitule": _texte(elem, "INTITULE") or "",
                "abrege_code": abrege_code,
                "abrege_lib": abrege_lib,
                "actif": actif_txt.startswith("o"),  # "Oui"
                "etat": _texte(elem, "ETAT_FICHE") or "",
                "date_fin": (parse_date_souple(_texte(elem, "DATE_FIN_ENREGISTREMENT")) or None),
                "successeurs": _codes_lies(elem, "NOUVELLE_CERTIFICATION"),
                "formacodes": _formacodes(elem),
            }
            n += 1
            if n % 5000 == 0:
                log(f"  … {n} fiches lues")
        elem.clear()
        racine.clear()
    if debug_fiche:
        raise SystemExit(f"Fiche {debug_fiche} introuvable dans l'export.")
    log(f"  {n} fiches RNCP chargées ({sum(1 for f in repertoire.values() if f['actif'])} actives)")
    if n == 0:
        raise SystemExit("Export vide ou schéma inattendu (aucun élément FICHE).")
    return repertoire


def successeurs_terminaux(code: str, repertoire: dict, profondeur_max: int = 8) -> list[str]:
    """Suit la chaîne de remplacement jusqu'aux fiches actives (ou terminales)."""
    resultat: list[str] = []
    vus: set[str] = set()
    frontiere = [(s, 1) for s in repertoire.get(code, {}).get("successeurs", [])]
    while frontiere:
        c, prof = frontiere.pop(0)
        if c in vus or prof > profondeur_max:
            continue
        vus.add(c)
        fiche = repertoire.get(c)
        if fiche is None or fiche["actif"] or not fiche["successeurs"]:
            if c not in resultat:
                resultat.append(c)
        else:
            frontiere.extend((s, prof + 1) for s in fiche["successeurs"])
    return resultat


# --------------------------------------------------------------------------- #
# 3. Lecture du fichier Excel (liste des codes suivis)
# --------------------------------------------------------------------------- #

def localiser_entetes(ws) -> tuple[int, dict[str, int]]:
    """Trouve la ligne d'en-tête et les index de colonnes utiles."""
    voulu = {"nomenclature rncp": "nomenclature", "code rncp": "code",
             "certification": "certification", "formacode": "formacode"}
    for ligne in range(1, 16):
        colonnes: dict[str, int] = {}
        for col in range(1, min(ws.max_column, 40) + 1):
            v = ws.cell(row=ligne, column=col).value
            if isinstance(v, str):
                cle = voulu.get(v.strip().casefold())
                if cle and cle not in colonnes:
                    colonnes[cle] = col
        if "nomenclature" in colonnes or "code" in colonnes:
            return ligne, colonnes
    raise SystemExit(
        "Impossible de trouver la ligne d'en-tête (colonnes « Nomenclature RNCP » "
        "ou « Code RNCP ») dans l'onglet principal."
    )


def lire_watchlist(chemin_excel: Path) -> dict:
    """Lit le classeur en valeurs (data_only) : codes suivis + intitulés par ligne."""
    wb = load_workbook(chemin_excel, data_only=True)
    nom_master = "Formacode_RNCP" if "Formacode_RNCP" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[nom_master]
    ligne_entete, colonnes = localiser_entetes(ws)
    col_code = colonnes.get("nomenclature") or colonnes.get("code")
    col_intitule = colonnes.get("certification")
    lignes: list[tuple[int, str, str]] = []   # (n° ligne, code, intitulé fichier)
    for r in range(ligne_entete + 1, ws.max_row + 1):
        code = canon(ws.cell(row=r, column=col_code).value)
        if code is None and colonnes.get("code"):
            code = canon(ws.cell(row=r, column=colonnes["code"]).value)
        if code is None:
            continue
        intitule = ws.cell(row=r, column=col_intitule).value if col_intitule else ""
        lignes.append((r, code, str(intitule or "").strip()))
    intitules: dict[str, str] = {}
    for _, code, intitule in lignes:
        if intitule and code not in intitules:
            intitules[code] = intitule
    wb.close()
    if not lignes:
        raise SystemExit("Aucun code RNCP trouvé dans l'onglet principal.")
    return {
        "master": nom_master,
        "ligne_entete": ligne_entete,
        "colonnes": colonnes,
        "lignes": lignes,
        "intitules": intitules,
        "codes": sorted(intitules.keys() | {c for _, c, _ in lignes}, key=lambda c: int(c[4:])),
    }


# --------------------------------------------------------------------------- #
# 4. Évaluation d'un code
# --------------------------------------------------------------------------- #

def _affiche(code: str, repertoire: dict) -> str:
    f = repertoire.get(code)
    if f and f["intitule"]:
        etat = "" if f["actif"] else " — fiche inactive"
        return f"{code} ({f['intitule']}{etat})"
    return code


def evaluer(code: str, intitule_excel: str, repertoire: dict, aujourdhui: date,
            seuil_jours: int) -> dict:
    fiche = repertoire.get(code)
    if fiche is None:
        return {
            "statut": "GRIS",
            "details": ["Code introuvable dans le répertoire RNCP — saisie à vérifier"],
            "codes_a_jour": "", "intitule_a_jour": "", "echeance": "",
            "formacodes_a_jour": "",
        }
    terminaux = successeurs_terminaux(code, repertoire)
    cible = terminaux[0] if terminaux else code
    fiche_cible = repertoire.get(cible, fiche)
    details: list[str] = []
    if not fiche["actif"]:
        statut = "ROUGE"
        if terminaux:
            details.append("Fiche inactive — remplacée par "
                           + " ; ".join(_affiche(c, repertoire) for c in terminaux))
        else:
            quand = f" (échéance {fr_date(fiche['date_fin'])})" if fiche["date_fin"] else ""
            details.append(f"Fiche inactive{quand} — aucun successeur publié")
    else:
        if fiche["successeurs"]:
            details.append("Successeur déjà publié : "
                           + " ; ".join(_affiche(c, repertoire) for c in terminaux)
                           + " — préparer la bascule de l'offre CPF")
        if fiche["date_fin"]:
            restant = (fiche["date_fin"] - aujourdhui).days
            if restant < 0:
                details.append(f"Échéance {fr_date(fiche['date_fin'])} dépassée "
                               "mais fiche encore marquée active")
            elif restant <= seuil_jours:
                details.append(f"Échéance le {fr_date(fiche['date_fin'])} (dans {restant} j)")
        candidats = {norm_titre(fiche["intitule"])}
        if fiche["abrege_code"]:
            candidats.add(norm_titre(f"{fiche['abrege_code']} - {fiche['intitule']}"))
        if fiche["abrege_lib"]:
            candidats.add(norm_titre(f"{fiche['abrege_lib']} - {fiche['intitule']}"))
        if intitule_excel and norm_titre(intitule_excel) not in candidats:
            details.append(f"Intitulé officiel différent : « {fiche['intitule']} »")
        statut = "ORANGE" if details else "VERT"
    return {
        "statut": statut,
        "details": details,
        "codes_a_jour": ", ".join(terminaux) if terminaux else code,
        "intitule_a_jour": fiche_cible["intitule"],
        "echeance": fr_date(fiche["date_fin"]),
        "formacodes_a_jour": " ; ".join(
            f"{c} ({l})" if l else c for c, l in fiche_cible["formacodes"]
        ),
    }


# --------------------------------------------------------------------------- #
# 5. Instantané et détection des changements
# --------------------------------------------------------------------------- #

def etat_pour_snapshot(code: str, ev: dict, repertoire: dict) -> dict:
    fiche = repertoire.get(code) or {}
    return {
        "statut": ev["statut"],
        "actif": bool(fiche.get("actif")) if fiche else None,
        "intitule": fiche.get("intitule", ""),
        "date_fin": fiche["date_fin"].isoformat() if fiche.get("date_fin") else "",
        "successeurs": fiche.get("successeurs", []),
    }


def comparer(ancien: dict, nouveau: dict, jour: str) -> list[dict]:
    events: list[dict] = []

    def ev(code, type_, avant, apres):
        events.append({"date": jour, "code": code, "type": type_,
                       "avant": str(avant), "apres": str(apres)})

    anciens, nouveaux = set(ancien), set(nouveau)
    for c in sorted(nouveaux - anciens, key=lambda x: int(x[4:])):
        ev(c, "Nouveau code suivi", "", LIBELLES[nouveau[c]["statut"]])
    for c in sorted(anciens - nouveaux, key=lambda x: int(x[4:])):
        ev(c, "Code retiré du fichier", LIBELLES.get(ancien[c].get("statut", ""), ""), "")
    for c in sorted(anciens & nouveaux, key=lambda x: int(x[4:])):
        o, n = ancien[c], nouveau[c]
        if o.get("statut") != n.get("statut"):
            ev(c, "Changement de statut",
               LIBELLES.get(o.get("statut", ""), o.get("statut", "")),
               LIBELLES.get(n.get("statut", ""), n.get("statut", "")))
        if o.get("actif") != n.get("actif"):
            ev(c, "État de la fiche", "Active" if o.get("actif") else "Inactive",
               "Active" if n.get("actif") else "Inactive")
        if o.get("intitule") != n.get("intitule"):
            ev(c, "Intitulé officiel modifié", o.get("intitule", ""), n.get("intitule", ""))
        if o.get("date_fin") != n.get("date_fin"):
            ev(c, "Échéance modifiée",
               fr_date(parse_date_souple(o.get("date_fin"))),
               fr_date(parse_date_souple(n.get("date_fin"))))
        if o.get("successeurs") != n.get("successeurs"):
            ev(c, "Successeur(s) modifié(s)",
               ", ".join(o.get("successeurs", [])), ", ".join(n.get("successeurs", [])))
    return events


# --------------------------------------------------------------------------- #
# 6. Production du classeur annoté
# --------------------------------------------------------------------------- #

def annoter_classeur(chemin_source: Path, chemin_sortie: Path, watch: dict,
                     evaluations: dict[str, dict], historique: list[dict],
                     meta: dict, events: list[dict], n_nouveaux: int) -> None:
    # Récapitulatif « ce qui vient de changer », par code
    changes_par_code: dict[str, list[str]] = {}
    for e in events:
        if not e["code"].startswith("RNCP"):
            continue
        if e["avant"] and e["apres"]:
            libelle = f"{e['type']} : {e['avant']} → {e['apres']}"
        elif e["avant"] or e["apres"]:
            libelle = f"{e['type']} : {e['avant'] or e['apres']}"
        else:
            libelle = e["type"]
        changes_par_code.setdefault(e["code"], []).append(libelle)

    wb = load_workbook(chemin_source)  # formules conservées
    try:
        wb.calculation.fullCalcOnLoad = True  # Excel recalcule à l'ouverture
    except Exception:  # noqa: BLE001
        pass

    ws = wb[watch["master"]]
    ligne_entete = watch["ligne_entete"]
    premiere_nouvelle = ws.max_column + 1

    # En-têtes des nouvelles colonnes
    for i, titre in enumerate(NOUVELLES_COLONNES):
        cell = ws.cell(row=ligne_entete, column=premiere_nouvelle + i, value=titre)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    largeurs = [16, 60, 22, 50, 14, 55, 10, 45]
    for i, l in enumerate(largeurs):
        ws.column_dimensions[get_column_letter(premiere_nouvelle + i)].width = l

    # Lignes de données : valeurs + couleur
    derniere_col = premiere_nouvelle + len(NOUVELLES_COLONNES) - 1
    col_changement = derniere_col          # détail du changement
    col_change_bool = derniere_col - 1     # « Changé ? » OUI/NON
    centre = Alignment(horizontal="center")
    for r, code, _ in watch["lignes"]:
        ev = evaluations[code]
        a_change = code in changes_par_code
        valeurs = [
            LIBELLES[ev["statut"]],
            " | ".join(ev["details"]),
            ev["codes_a_jour"],
            ev["intitule_a_jour"],
            ev["echeance"],
            ev["formacodes_a_jour"],
            "OUI" if a_change else "NON",
            " ; ".join(changes_par_code.get(code, [])),
        ]
        for i, v in enumerate(valeurs):
            ws.cell(row=r, column=premiere_nouvelle + i, value=v)
        fill = FILLS[ev["statut"]]
        for col in range(1, derniere_col + 1):
            ws.cell(row=r, column=col).fill = fill
        ws.cell(row=r, column=premiere_nouvelle).font = FONTS_STATUT[ev["statut"]]
        ws.cell(row=r, column=col_change_bool).alignment = centre
        if a_change:  # nouveauté de ce contrôle : cellules bleues
            for cc in (col_change_bool, col_changement):
                c = ws.cell(row=r, column=cc)
                c.fill = FILL_NOUVEAU
                c.font = FONT_NOUVEAU

    # Filtres sur l'onglet maître : le fichier d'origine contient un tableau
    # Excel (« Tableau1 », B4:G5004) qui possède déjà ses boutons de filtre.
    # Superposer un filtre de feuille rendrait le classeur invalide (Excel
    # propose alors une « réparation ») : on étend donc le tableau existant
    # aux nouvelles colonnes. À défaut de tableau, filtre de feuille classique.
    table_maitre = None
    for t in ws.tables.values():
        try:
            min_c, min_r, max_c, max_r = range_boundaries(t.ref)
        except Exception:  # noqa: BLE001
            continue
        if min_r == ligne_entete:
            table_maitre = t
            break
    if table_maitre is not None:
        min_c, min_r, max_c, max_r = range_boundaries(table_maitre.ref)
        nouveau_ref = (f"{get_column_letter(min_c)}{min_r}:"
                       f"{get_column_letter(derniere_col)}{max(max_r, ws.max_row)}")
        noms_existants = {c.name for c in table_maitre.tableColumns}
        prochain_id = max((int(c.id) for c in table_maitre.tableColumns), default=0) + 1
        for titre in NOUVELLES_COLONNES:
            if titre in noms_existants:
                continue
            table_maitre.tableColumns.append(TableColumn(id=prochain_id, name=titre))
            prochain_id += 1
        table_maitre.ref = nouveau_ref
        if table_maitre.autoFilter is not None:
            table_maitre.autoFilter.ref = nouveau_ref
        else:
            table_maitre.autoFilter = AutoFilter(ref=nouveau_ref)
    else:
        ws.auto_filter.ref = (
            f"A{ligne_entete}:{get_column_letter(derniere_col)}{ws.max_row}"
        )

    # Onglets par formacode : coloration du code + de l'intitulé
    for feuille in wb.worksheets:
        if feuille.title in (watch["master"], "Synthèse", "Changements"):
            continue
        for ligne in feuille.iter_rows(min_col=1, max_col=1):
            cellule = ligne[0]
            code = canon(cellule.value if not str(cellule.value or "").startswith("=") else None)
            if code and code in evaluations:
                fill = FILLS[evaluations[code]["statut"]]
                cellule.fill = fill
                feuille.cell(row=cellule.row, column=2).fill = fill

    # ---- Onglet Synthèse ----
    for nom in ("Synthèse", "Changements"):
        if nom in wb.sheetnames:
            del wb[nom]
    syn = wb.create_sheet("Synthèse", 0)
    gras = Font(bold=True)
    syn["A1"] = f"Contrôle RNCP — {meta['date_controle']}"
    syn["A1"].font = Font(bold=True, size=14)
    syn["A2"] = f"Export France compétences : {meta['source_export']}"
    syn["A3"] = (f"{meta['nb_codes']} codes RNCP distincts contrôlés "
                 f"({meta['nb_lignes']} lignes du fichier)")
    ligne = 5
    syn.cell(row=ligne, column=1, value="Statut").font = gras
    syn.cell(row=ligne, column=2, value="Codes").font = gras
    for statut in STATUTS:
        ligne += 1
        c1 = syn.cell(row=ligne, column=1, value=LIBELLES[statut])
        c1.fill = FILLS[statut]
        c1.font = FONTS_STATUT[statut]
        syn.cell(row=ligne, column=2, value=meta["compteurs"].get(statut, 0))
    ligne += 2
    syn.cell(row=ligne, column=1,
             value="Codes nécessitant une action ou une vérification "
                   "(● = nouveauté de ce contrôle) :").font = gras
    ligne += 1
    entetes = ["Nouveau", "Code RNCP", "Intitulé (fichier)", "Statut", "Détail",
               "Code(s) à jour", "Intitulé à jour", "Formacodes de la fiche à jour",
               "Échéance fiche"]
    ligne_actions = ligne
    for i, t in enumerate(entetes, start=1):
        c = syn.cell(row=ligne, column=i, value=t)
        c.font = gras
        c.fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    syn.freeze_panes = syn.cell(row=ligne + 1, column=1)
    ordre = {"ROUGE": 0, "ORANGE": 1, "GRIS": 2}
    a_lister = [(c, evaluations[c]) for c in watch["codes"]
                if evaluations[c]["statut"] != "VERT"]
    a_lister.sort(key=lambda x: (ordre.get(x[1]["statut"], 9),
                                 0 if x[0] in changes_par_code else 1,
                                 int(x[0][4:])))
    for code, ev in a_lister:
        ligne += 1
        nouveau = code in changes_par_code
        valeurs = ["●" if nouveau else "", code, watch["intitules"].get(code, ""),
                   LIBELLES[ev["statut"]], " | ".join(ev["details"]),
                   ev["codes_a_jour"], ev["intitule_a_jour"],
                   ev["formacodes_a_jour"], ev["echeance"]]
        for i, v in enumerate(valeurs, start=1):
            syn.cell(row=ligne, column=i, value=v)
        if nouveau:
            syn.cell(row=ligne, column=1).fill = FILL_NOUVEAU
            syn.cell(row=ligne, column=1).font = FONT_NOUVEAU
        syn.cell(row=ligne, column=4).fill = FILLS[ev["statut"]]
        syn.cell(row=ligne, column=4).font = FONTS_STATUT[ev["statut"]]
    for col, largeur in zip("ABCDEFGHI", (10, 13, 55, 16, 75, 20, 50, 50, 13)):
        syn.column_dimensions[col].width = largeur
    syn.auto_filter.ref = f"A{ligne_actions}:I{max(ligne, ligne_actions)}"

    # ---- Onglet Changements (journal) ----
    chg = wb.create_sheet("Changements", 1)
    titre = chg.cell(row=1, column=1,
                     value="Journal des changements — les lignes surlignées en bleu "
                           "sont les nouveautés de ce contrôle (contenu du dernier envoi)")
    titre.font = Font(bold=True, color="1F4E79")
    for i, t in enumerate(["Date", "Code RNCP", "Changement", "Avant", "Après"], start=1):
        c = chg.cell(row=2, column=i, value=t)
        c.font = gras
        c.fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    chg.freeze_panes = "A3"
    for j, e in enumerate(reversed(historique), start=3):  # plus récent en haut
        for i, v in enumerate([e["date"], e["code"], e["type"], e["avant"], e["apres"]],
                              start=1):
            cellule = chg.cell(row=j, column=i, value=v)
            if j - 3 < n_nouveaux:  # lignes du contrôle courant
                cellule.fill = FILL_NOUVEAU
    for col, largeur in zip("ABCDE", (12, 13, 26, 55, 55)):
        chg.column_dimensions[col].width = largeur
    chg.auto_filter.ref = f"A2:E{max(3, 2 + len(historique))}"

    wb.save(chemin_sortie)


# --------------------------------------------------------------------------- #
# 7. Résumés (e-mail / journal d'exécution)
# --------------------------------------------------------------------------- #

def construire_resumes(meta: dict, events: list[dict], evaluations: dict,
                       watch: dict, premiere_execution: bool) -> tuple[str, str, str]:
    """-> (markdown, html, sujet)"""
    cpt = meta["compteurs"]
    bilan = (f"{cpt.get('ROUGE', 0)} rouge(s), {cpt.get('ORANGE', 0)} orange(s), "
             f"{cpt.get('GRIS', 0)} introuvable(s), {cpt.get('VERT', 0)} OK")
    if premiere_execution:
        sujet = f"Initialisation du suivi — {bilan}"
    else:
        sujet = f"{len(events)} changement(s) — {bilan}"

    lignes_md = [f"# Contrôle RNCP du {meta['date_controle']}", "",
                 f"Export France compétences : {meta['source_export']}  ",
                 f"Codes contrôlés : {meta['nb_codes']}  ", f"Bilan : **{bilan}**", ""]
    if premiere_execution:
        lignes_md.append("Première exécution : l'état des lieux complet est dans le "
                         "classeur joint (onglet Synthèse).")
    elif events:
        lignes_md.append("## Changements depuis le dernier contrôle")
        lignes_md.append("")
        lignes_md.append("| Code | Changement | Avant | Après |")
        lignes_md.append("|---|---|---|---|")
        for e in events:
            lignes_md.append(f"| {e['code']} | {e['type']} | {e['avant']} | {e['apres']} |")
        lignes_md.append("")

    ordre = {"ROUGE": 0, "ORANGE": 1}
    actions = [(c, evaluations[c]) for c in watch["codes"]
               if evaluations[c]["statut"] in ("ROUGE", "ORANGE")]
    actions.sort(key=lambda x: (ordre[x[1]["statut"]], int(x[0][4:])))
    if actions:
        lignes_md.append("## Codes à traiter")
        lignes_md.append("")
        for code, ev in actions[:60]:
            lignes_md.append(f"- **{code}** ({LIBELLES[ev['statut']]}) — "
                             f"{watch['intitules'].get(code, '')} : "
                             f"{' | '.join(ev['details'])}")
        if len(actions) > 60:
            lignes_md.append(f"- … et {len(actions) - 60} autres (voir classeur joint)")
    md = "\n".join(lignes_md) + "\n"

    # Version HTML compacte
    def chip(statut):
        fond, texte = COULEURS[statut]
        return (f'<span style="background:#{fond};color:#{texte};padding:1px 8px;'
                f'border-radius:9px;font-weight:600">{LIBELLES[statut]}</span>')

    h = [f"<h2>Contrôle RNCP du {meta['date_controle']}</h2>",
         f"<p>Export France compétences : {meta['source_export']}<br>"
         f"Codes contrôlés : {meta['nb_codes']}<br>"
         + " · ".join(f"{chip(s)} {cpt.get(s, 0)}" for s in STATUTS) + "</p>"]
    if premiere_execution:
        h.append("<p>Première exécution : état des lieux complet dans le classeur joint "
                 "(onglet <b>Synthèse</b>).</p>")
    elif events:
        h.append("<h3>Changements depuis le dernier contrôle</h3>")
        h.append('<table border="1" cellspacing="0" cellpadding="5" '
                 'style="border-collapse:collapse;font-family:Arial;font-size:13px">')
        h.append("<tr style='background:#DDEBF7'><th>Code</th><th>Changement</th>"
                 "<th>Avant</th><th>Après</th></tr>")
        for e in events:
            h.append(f"<tr><td>{e['code']}</td><td>{e['type']}</td>"
                     f"<td>{e['avant']}</td><td>{e['apres']}</td></tr>")
        h.append("</table>")
    if actions:
        h.append("<h3>Codes à traiter</h3><ul>")
        for code, ev in actions[:60]:
            h.append(f"<li><b>{code}</b> {chip(ev['statut'])} — "
                     f"{watch['intitules'].get(code, '')}<br>"
                     f"<small>{' | '.join(ev['details'])}</small></li>")
        if len(actions) > 60:
            h.append(f"<li>… et {len(actions) - 60} autres (voir classeur joint)</li>")
        h.append("</ul>")
    h.append("<p><small>Détail complet : classeur joint (onglets Synthèse, Changements "
             "et Formacode_RNCP annoté).</small></p>")
    html = "\n".join(h) + "\n"
    return md, html, sujet


def ecrire_github_output(changes: bool, sujet: str, resume: str) -> None:
    import os
    chemin = os.environ.get("GITHUB_OUTPUT")
    if not chemin:
        return
    sujet = re.sub(r"[\r\n]+", " ", sujet)
    resume = re.sub(r"[\r\n\"'`]+", " ", resume)
    with open(chemin, "a", encoding="utf-8") as f:
        f.write(f"changes={'true' if changes else 'false'}\n")
        f.write(f"sujet={sujet}\n")
        f.write(f"resume={resume}\n")


# --------------------------------------------------------------------------- #
# Programme principal
# --------------------------------------------------------------------------- #

def main() -> None:
    p = argparse.ArgumentParser(description="Contrôle des fiches RNCP (France compétences)")
    p.add_argument("--excel", default="", help="classeur à contrôler (défaut : data/*.xlsx)")
    p.add_argument("--state", default="state/snapshot.json")
    p.add_argument("--outdir", default="output")
    p.add_argument("--xml", default="", help="export XML local (ne pas télécharger)")
    p.add_argument("--seuil-jours", type=int, default=180,
                   help="alerte orange si l'échéance est à moins de N jours (défaut 180)")
    p.add_argument("--min-couverture", type=float, default=0.5,
                   help="part minimale des codes du fichier retrouvés dans l'export")
    p.add_argument("--debug-fiche", default="",
                   help="affiche le XML brut d'une fiche (ex. RNCP35803) puis s'arrête")
    args = p.parse_args()

    chemin_excel = Path(args.excel) if args.excel else None
    if chemin_excel is None or not chemin_excel.exists():
        candidats = sorted(Path("data").glob("*.xlsx"))
        if not candidats:
            raise SystemExit("Aucun classeur trouvé (option --excel ou dossier data/).")
        chemin_excel = candidats[0]

    try:
        from zoneinfo import ZoneInfo
        maintenant = datetime.now(ZoneInfo("Europe/Paris"))
    except Exception:  # noqa: BLE001
        maintenant = datetime.now(timezone.utc)
    aujourdhui = maintenant.date()
    jour_iso = aujourdhui.isoformat()

    log(f"Fichier contrôlé : {chemin_excel}")
    watch = lire_watchlist(chemin_excel)
    log(f"  {len(watch['codes'])} codes RNCP distincts sur {len(watch['lignes'])} lignes")

    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = Path(tmp)
        if args.xml:
            xml_path, source = Path(args.xml), Path(args.xml).name
            log(f"Export local : {xml_path}")
        else:
            log("Recherche du dernier export RNCP sur data.gouv.fr…")
            with requests.Session() as session:
                xml_path, source = telecharger_et_extraire(session, tmpdir)
        log("Lecture de l'export…")
        repertoire = parser_export(xml_path, args.debug_fiche or None)

    trouves = sum(1 for c in watch["codes"] if c in repertoire)
    couverture = trouves / max(len(watch["codes"]), 1)
    log(f"  Codes du fichier retrouvés dans l'export : {trouves}/{len(watch['codes'])} "
        f"({couverture:.0%})")
    if couverture < args.min_couverture:
        raise SystemExit(
            "Couverture anormalement basse : le schéma de l'export a peut-être changé. "
            "Lancer :  python check_rncp.py --debug-fiche RNCP35803  pour vérifier."
        )

    evaluations = {c: evaluer(c, watch["intitules"].get(c, ""), repertoire,
                              aujourdhui, args.seuil_jours)
                   for c in watch["codes"]}
    compteurs = {s: sum(1 for e in evaluations.values() if e["statut"] == s)
                 for s in STATUTS}
    log("Bilan : " + ", ".join(f"{compteurs[s]} {LIBELLES[s]}" for s in STATUTS))

    # Instantané précédent -> détection des changements
    chemin_state = Path(args.state)
    premiere_execution = not chemin_state.exists()
    snapshot_prec = {}
    historique: list[dict] = []
    if not premiere_execution:
        contenu = json.loads(chemin_state.read_text(encoding="utf-8"))
        snapshot_prec = contenu.get("codes", {})
        historique = contenu.get("historique", [])
    etat_actuel = {c: etat_pour_snapshot(c, evaluations[c], repertoire)
                   for c in watch["codes"]}
    if premiere_execution:
        events: list[dict] = []
        historique.append({"date": jour_iso, "code": "—",
                           "type": "Initialisation du suivi",
                           "avant": "", "apres": f"{len(watch['codes'])} codes suivis"})
    else:
        events = comparer(snapshot_prec, etat_actuel, jour_iso)
        historique.extend(events)
    historique = historique[-1000:]

    meta = {
        "date_controle": maintenant.strftime("%d/%m/%Y %H:%M"),
        "source_export": source,
        "nb_codes": len(watch["codes"]),
        "nb_lignes": len(watch["lignes"]),
        "compteurs": compteurs,
    }

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    sortie_xlsx = outdir / "Liens_CPF_controle.xlsx"
    log(f"Écriture du classeur annoté : {sortie_xlsx}")
    n_nouveaux = 1 if premiere_execution else len(events)
    annoter_classeur(chemin_excel, sortie_xlsx, watch, evaluations, historique, meta,
                     events, n_nouveaux)

    md, html, sujet = construire_resumes(meta, events, evaluations, watch,
                                         premiere_execution)
    (outdir / "summary.md").write_text(md, encoding="utf-8")
    (outdir / "summary.html").write_text(html, encoding="utf-8")

    chemin_state.parent.mkdir(parents=True, exist_ok=True)
    chemin_state.write_text(
        json.dumps({"derniere_execution": maintenant.isoformat(),
                    "source_export": source, "codes": etat_actuel,
                    "historique": historique},
                   ensure_ascii=False, indent=1),
        encoding="utf-8",
    )

    changements = premiere_execution or bool(events)
    resume = (f"{len(events)} changement(s), {compteurs['ROUGE']} rouge(s), "
              f"{compteurs['ORANGE']} orange(s)")
    if premiere_execution:
        resume = "initialisation, " + resume
    ecrire_github_output(changements, sujet, resume)

    import os
    step_summary = os.environ.get("GITHUB_STEP_SUMMARY")
    if step_summary:
        with open(step_summary, "a", encoding="utf-8") as f:
            f.write(md)

    log(f"Changements à signaler : {'oui' if changements else 'non'} — {sujet}")


if __name__ == "__main__":
    main()
